#!/usr/bin/env python3
from docx import Document
doc = Document(r'C:\Users\lenovo\Desktop\3024244171_牛泽钰_第三周课程报告.docx')

fixes = {
    '本阶段新增以下拥塞控制变量，均定义在struct tcp_cb中（tju_tcp.c）：':
        '拥塞控制相关变量定义在struct tcp_cb中（tju_tcp.c:110-113）：',
    '由于慢启动阶段每RTT内cwnd翻倍，因此也称为指数增长。':
        'cwnd每RTT翻倍。',
    '选择以下两类因素开展对照实验：\n因素一：丢包率（0% vs 10%），固定参数：带宽100Mbps、延迟300ms、延迟波动50ms。每组实验重复1次（每次60秒数据传输），平均吞吐率由gen_graph_win.py计算（总接收字节数/总传输时间）。\n因素二：网络延迟（50ms vs 300ms），固定参数：带宽100Mbps、0%丢包、0延迟波动。\n所有实验使用相同代码版本（USE_CWND_LIMIT=1，INIT_WDSIZE=10，RING_BYTES=7MB），测试脚本为test_congestion.py，网络条件通过tc命令设置：带宽、延迟、延迟分布和丢包率。':
        '对照实验从两个维度展开：\n丢包率维度（0% vs 10%）：固定带宽100Mbps、延迟300ms、延迟波动50ms，每组60秒数据传输，吞吐率由gen_graph_win.py统计（总接收字节/总传输时间）。\n网络延迟维度（50ms vs 300ms）：固定带宽100Mbps、0%丢包。\n所有实验使用同一代码版本（USE_CWND_LIMIT=1，INIT_WDSIZE=10，RING_BYTES=7MB），网络条件由test_congestion.py通过tc命令设置。',
}

count = 0
for p in doc.paragraphs:
    for old, new in fixes.items():
        if p.text.strip()[:30] == old[:30]:
            for run in p.runs:
                run.text = ''
            if p.runs:
                p.runs[0].text = new
            else:
                p.add_run(new)
            count += 1
            break

doc.save(r'C:\Users\lenovo\Desktop\3024244171_牛泽钰_第三周课程报告.docx')
print(f"Fixed {count} paragraphs")
