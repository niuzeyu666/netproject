#!/usr/bin/env python3
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import copy

doc = Document(r'C:\Users\lenovo\Desktop\3024244171_牛泽钰_第三周课程报告.docx')

def find_para(text_start):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(text_start):
            return i
    return None

def insert_image_after(anchor_idx, img_path, caption, width_inches=5.5):
    anchor_p = doc.paragraphs[anchor_idx]
    img_p = copy.deepcopy(anchor_p._p)
    anchor_p._p.addnext(img_p)
    img_para = None
    for p in doc.paragraphs:
        if p._p is img_p:
            img_para = p
            break
    for run in img_para.runs:
        run.text = ''
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_para.add_run()
    run.add_picture(img_path, width=Inches(width_inches))
    cap_p = copy.deepcopy(anchor_p._p)
    img_p.addnext(cap_p)
    for p in doc.paragraphs:
        if p._p is cap_p:
            p.text = caption
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

# Find fig1 caption and insert fig3 after it
idx = find_para('图1 10%丢包')
if idx:
    insert_image_after(idx,
        r'C:\netproject\tju_tcp\figure\smallrwnd\AllWindowSize_VS_Time_0.044273_8.675225.png',
        '图3 小rwnd场景下cwnd（红）、rwnd（绿）、swnd（蓝）三窗口对比，swnd被rwnd钳制')
    print(f"Fig3 inserted after fig1 caption at para {idx}")
else:
    print("Fig1 caption not found!")

doc.save(r'C:\Users\lenovo\Desktop\3024244171_牛泽钰_第三周课程报告.docx')
print("Done!")
