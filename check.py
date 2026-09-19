from docx import Document
doc = Document(r"C:\Users\lenovo\Desktop\3024244171_牛泽钰_第三周课程报告.docx")
issues = []
for i, p in enumerate(doc.paragraphs):
    if "【待" in p.text or "【填写" in p.text or "【仅在" in p.text:
        issues.append(f"P{i}: {p.text[:60]}")
for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            if "【待" in cell.text or "【填写" in cell.text or "【仅在" in cell.text:
                issues.append(f"T{ti}R{ri}C{ci}: {cell.text[:60]}")
print(f"Remaining placeholders: {len(issues)}")
for iss in issues:
    print(f"  {iss}")
