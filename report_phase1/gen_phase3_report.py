from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\netproject\reports\3024244171_牛泽钰_第三阶段实验报告.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.3)
sec.bottom_margin = Cm(2.3)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.6)

styles = doc.styles
styles["Normal"].font.name = "宋体"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(10.5)
for name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 13)]:
    styles[name].font.name = "黑体"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

def title(text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)

def h(text, level=1):
    doc.add_heading(text, level=level)

def para(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    return p

def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.line_spacing = 1.3

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, text in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = text
        shade(c, "D9EAF7")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in c.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph()
    return t

title("计算机网络实践第三阶段实验报告")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("TCP 拥塞控制实现与验证\n").bold = True
p.add_run("学号 3024244171    姓名 牛泽钰    日期 2026年9月19日")

h("一 实验目标")
para("本阶段在第二阶段可靠传输实现之上，依据 RFC 5681 实现 Reno/NewReno 拥塞控制。发送端使用拥塞窗口 cwnd 与接收端通告窗口 rwnd 的较小值约束在途数据，完成慢启动、拥塞避免、快速重传、快速恢复以及超时后的窗口回退，并保留可复现的事件 trace。")

h("二 实现边界")
table(["范围", "内容"], [
    ["本阶段实现", "Slow Start、Congestion Avoidance、Fast Retransmit、Fast Recovery、RTO 回退、FlightSize 约束"],
    ["第二阶段复用", "连接建立与关闭、累计 ACK、乱序缓存、RTT/RTO、接收窗口"],
    ["未纳入", "CUBIC、SACK、ECN、BBR、RACK-TLP 等挑战功能"],
], [3.2, 12.5])

h("三 数据结构与发送约束")
table(["字段", "单位", "作用"], [
    ["cwnd", "字节", "拥塞窗口"], ["ssthresh", "字节", "慢启动阈值"],
    ["peer_window", "字节", "接收端通告窗口"], ["tx_nxt - tx_una", "字节", "FlightSize"],
    ["dup_ack_cnt", "次", "重复 ACK 计数"], ["recover_point", "序号", "快速恢复退出边界"],
], [3.2, 2.4, 10.2])
para("发送函数 flush_tx 计算 effective_window=min(cwnd,rwnd)，仅当 effective_window-FlightSize 大于零时发送新数据。INIT_WDS 可由 Makefile 传入并转换为 INIT_WDSIZE 个 SMSS 字节。")

h("四 RFC 5681 状态转换")
h("4.1 慢启动", 2)
para("当 cwnd 小于 ssthresh 时，每个确认新数据的累计 ACK 使 cwnd 增长 min(newly_acked,SMSS)。窗口随往返轮次近似指数增长。")
h("4.2 拥塞避免", 2)
para("当 cwnd 不小于 ssthresh 时，按 SMSS*newly_acked/cwnd 增长，使一个 RTT 内累计增长约一个 SMSS。所有运算使用字节单位和 64 位中间值避免溢出。")
h("4.3 快速重传与快速恢复", 2)
para("对未推进 tx_una 且仍有在途数据的 ACK 计数。第三个重复 ACK 到达时，ssthresh=max(FlightSize/2,2*SMSS)，cwnd=ssthresh+3*SMSS，记录 recover_point 并立即重传最早未确认段。额外重复 ACK 使 cwnd 增加一个 SMSS；确认越过 recover_point 后退出快速恢复并令 cwnd=ssthresh。")
h("4.4 超时", 2)
para("RTO 超时视为严重拥塞：ssthresh=max(FlightSize/2,2*SMSS)，cwnd=SMSS，退出快速恢复并重传最早未确认段。RTO 指数退避，重传段不用于 RTT 采样，符合 Karn 算法。")

h("五 实现修正记录")
table(["问题", "影响", "修正"], [
    ["发送路径绕过 cwnd", "拥塞控制曲线无效", "启用 USE_CWND_LIMIT，发送窗口取 min(cwnd,rwnd)"],
    ["ssthresh 最低值为 49152", "不符合 RFC 减半下限", "改为 2*SMSS"],
    ["INIT_WDS 参数未生效", "测试不同初始窗口得到相同结果", "Makefile 同时支持 INIT_WDS 与 INIT_WDSIZE"],
    ["握手失败后过早释放 socket", "API 线程出现 use-after-free", "内核表摘除与应用所有权分离"],
    ["本地测试地址被改为 172.17.0.6", "与 client/server 的 .2/.3 拓扑不符", "恢复服务端地址 172.17.0.3"],
], [4.2, 5.2, 6.5])

h("六 测试环境与步骤")
table(["项目", "配置"], [
    ["拓扑", "client 172.17.0.2 / server 172.17.0.3"],
    ["编译器", "gcc -pthread -g -ggdb"],
    ["报文", "UDP 承载的 20 字节课程 TCP 头，MAX_SEG_LEN=1375"],
    ["可变参数", "rate、delay、delay-distro、loss、INIT_WDS"],
], [4.2, 11.5])
para("编译验证依次执行 make INIT_WDS=1、10、50。拥塞测试先启动 server 与 tcpdump，再启动 client；建立连接后用 tcset 设置带宽、时延、抖动和丢包率，运行规定时长，终止双端并恢复 100Mbps/20ms 网络。最后收集 client.event.trace、server.event.trace 与 server.pcap。")

h("七 测试矩阵")
table(["编号", "rate", "delay", "distro", "loss", "INIT_WDS", "检查点"], [
    ["CC-01", "100Mbps", "20ms", "0", "0%", "1", "慢启动"],
    ["CC-02", "100Mbps", "20ms", "0", "0%", "10", "慢启动转拥塞避免"],
    ["CC-03", "100Mbps", "300ms", "50", "5%", "10", "快速重传/恢复"],
    ["CC-04", "100Mbps", "300ms", "50", "10%", "10", "超时回退"],
    ["CC-05", "50Mbps", "300ms", "50", "5%", "50", "大初始窗口"],
], [1.4, 2.3, 2.1, 2.0, 1.8, 2.4, 5.3])

h("八 已完成验证")
bullet("INIT_WDS=1、10、50 三种配置均在课程 Ubuntu VM 使用 GCC 编译通过。")
bullet("编译命令已显示 -DINIT_WDSIZE 对应实际传入值，确认参数链路有效。")
bullet("发送窗口已恢复 RFC 要求的 min(cwnd,rwnd) 约束。")
bullet("拥塞事件输出 CWND、DUPACK、FAST_RETRANSMIT、TIMEOUT、ssthresh 与 recover_point 字段。")
bullet("修复连接失败路径的 socket 生命周期错误，避免控制块在 API 返回前释放。")

h("九 证据文件与复现")
para("源码位于 tju_tcp/src 与 tju_tcp/inc；构建入口为 tju_tcp/Makefile；测试驱动与绘图工具位于 tju_tcp/test。第三阶段原始证据统一放入 reports/phase3_evidence，包括编译日志、测试矩阵、trace、pcap 和生成图。提交前应从 client VM 执行完整测试矩阵，并用仓库内脚本重新生成窗口曲线。")

h("十 AI 协作记录")
para("AI 用于对照 RFC 5681 审核拥塞窗口更新、识别 Makefile 参数未传递、发现 cwnd 约束被关闭以及定位 socket 生命周期崩溃。所有修改均通过源码定位和 Linux VM 编译验证；报告未将未运行的极端网络测试描述为已通过。")

h("十一 结论")
para("第三阶段代码已建立符合 RFC 5681 边界的 Reno/NewReno 控制路径，并提供可配置初始窗口及结构化 trace。最终验收以完整测试矩阵生成的窗口曲线、抓包和第二阶段回归结果为准。")

doc.save(OUT)
print(OUT)
