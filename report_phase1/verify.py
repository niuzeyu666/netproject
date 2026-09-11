# -*- coding: utf-8 -*-
from docx import Document
doc = Document(r'C:\netproject\report_phase1\计算机网络实践课程报告-第一阶段-V3.docx')
checks = {
 '封面仓库地址': 'niuzeyu666/netproject' in doc.tables[0].rows[6].cells[1].text,
 '2.1标准依据差异': any('RFC 9293（TCP基本规范）' in p.text for p in doc.paragraphs),
 '2.2需求追踪表': doc.tables[3].rows[1].cells[1].text=='连接建立与关闭',
 '3.1环境表': 'Vagrant' in doc.tables[4].rows[1].cells[1].text,
 '3.2项目结构': any('项目目录结构' in p.text for p in doc.paragraphs),
 '3.3接口数据流': any('主要函数调用关系与数据流图' in p.text for p in doc.paragraphs),
 '3.4数据结构并发': any('关键数据结构分析' in p.text for p in doc.paragraphs),
 '3.5环境测试': any('环境搭建测试结果' in p.text for p in doc.paragraphs),
 '3.6AI记录': any('AI工具' in p.text and '双线程架构' in p.text for p in doc.paragraphs),
 '4.1设计目标': any('设计目标' in p.text for p in doc.paragraphs),
 '4.2报文格式': any('自定义20字节报文头' in p.text for p in doc.paragraphs),
 '4.3状态机窗口': any('TCP连接状态机FSM' in p.text for p in doc.paragraphs),
 '4.4并发异常': any('并发控制设计' in p.text for p in doc.paragraphs),
 '4.5分阶段表': '环境搭建与基线运行验证' in doc.tables[5].rows[1].cells[1].text,
 '4.6AI记录': any('发送窗口与接收窗口协同设计' in p.text for p in doc.paragraphs),
 '附录A错误修正': any('tju_recv阻塞机制的误判' in p.text for p in doc.paragraphs),
 '附录B进度表': '完成环境搭建' in doc.tables[10].rows[1].cells[1].text,
 '诚信声明日期': any('2026年09月06日' in p.text for p in doc.paragraphs),
}
passed=sum(1 for v in checks.values() if v)
for k,v in checks.items():
    s='PASS' if v else 'FAIL'
    print(f'  [{s}] {k}')
print(f'\n{passed}/{len(checks)} 通过')
total=sum(len(p.text) for p in doc.paragraphs)+sum(len(c.text) for t in doc.tables for r in t.rows for c in r.cells)
print(f'总字数: {total}, 段落: {len(doc.paragraphs)}, 表格: {len(doc.tables)}')
hints=[p for p in doc.paragraphs if '填写提示' in p.text and any(k in p.text for k in ['目录/模块','函数调用','tju_tcp_t','环境搭建测试','AI架构分析','总体架构图','20字节报文头','连接状态','线程/锁','AI工具与使用','对结果有实际影响','基线程序'])]
print(f'第一阶段遗留提示: {len(hints)}')
