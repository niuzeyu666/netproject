# -*- coding: utf-8 -*-
"""V3模板 - 附录A/B + 清理"""
from docx import Document
DST = r'C:\netproject\report_phase1\计算机网络实践课程报告-第一阶段-V3.docx'
doc = Document(DST)
def rp(p,t):
    if not p.runs: p.add_run(t); return
    p.runs[0].text=t
    for r in p.runs[1:]: r._element.getparent().remove(r._element)
def fp(k):
    for p in doc.paragraphs:
        if k in p.text: return p
    return None
def sc(c,t):
    for p in c.paragraphs:
        for r in p.runs: r.text=''
    if c.paragraphs[0].runs: c.paragraphs[0].runs[0].text=t
    else: c.paragraphs[0].add_run(t)

# 附录A
p=fp('至少选择一项对结果有实际影响的案例')
if p: rp(p,
 '【AI错误修正案例：tju_recv阻塞机制的误判】\n'
 '(1)AI错误描述：在框架架构分析过程中，AI最初判断tju_recv函数使用了pthread_cond_t条件变量(wait_cond)实现阻塞等待，认为并发模型设计合理。AI的理由是：tju_socket()中初始化了wait_cond(tju_tcp.c第20行调用pthread_cond_init)，global.h中也定义了该字段，因此推断tju_recv使用pthread_cond_wait阻塞。\n'
 '(2)错误发现过程：在进行人工核验时，逐行审查tju_recv函数(tju_tcp.c第137-167行)，发现第138行是while(sock->received_len<=0){}的纯忙等循环，没有任何pthread_cond_wait调用。进一步全文搜索pthread_cond_wait和pthread_cond_signal，在整个tju_tcp.c中均无匹配——wait_cond虽然在tju_socket中初始化(第20行)，但在整个代码库中从未被使用。\n'
 '(3)AI为何出错：AI基于"应该如此"的推断而非实际代码审查。框架中预留了wait_cond字段并进行了初始化，AI看到条件变量已初始化就推断它被使用了，没有逐行检查tju_recv的实际实现。这是AI代码分析的典型问题——容易将"已声明/已初始化"等同于"已使用"，将"设计意图"等同于"实际实现"。\n'
 '(4)对结果的实际影响：a)忙等会导致CPU占用率100%，在性能测试中会严重干扰吞吐率和CPU使用率指标；b)如果不修正，第二阶段实现时可能沿用忙等模式，导致并发性能问题和能源浪费；c)条件变量的正确使用是本项目并发控制的核心设计之一，必须在设计阶段明确。\n'
 '(5)如何发现：通过逐行代码审查(tju_tcp.c:137-167)+全文关键字搜索(pthread_cond_wait/pthread_cond_signal)+运行时观察(虚拟机CPU占用率接近100%)三重验证确认。\n'
 '(6)采用的证据：tju_tcp.c第138行while(sock->received_len<=0){}(忙等循环源代码)；tju_tcp.c第20行pthread_cond_init(&sock->wait_cond,NULL)(条件变量初始化但未使用)；全文搜索结果pthread_cond_wait和pthread_cond_signal在tju_tcp.c中出现0次；运行时top/htop输出进程CPU占用率接近100%。\n'
 '(7)最终修正：a)在3.4节并发机制分析中明确指出wait_cond已初始化但完全未使用，tju_recv通过忙等阻塞会导致CPU占用100%；b)在4.4节并发控制设计中明确规定将忙等改为pthread_cond_wait(&sock->wait_cond,&sock->recv_lock)，并在tju_handle_packet收到数据后调用pthread_cond_signal(&sock->wait_cond)唤醒；c)在3.6节AI协作记录中将此作为AI错误修正案例记录；d)第二阶段实现时严格按照修正后的设计执行。\n'
 '(8)经验总结：AI对代码的分析可能基于"应该如此"的推断而非实际代码审查，尤其是当框架中预留了某些字段(如条件变量)但未实际使用时，AI容易误判为已使用。因此所有AI分析结论必须经过逐行代码定位和实际运行验证，不能仅凭AI输出就采纳。本案例也体现了报告中"AI分析结论必须人工核验"要求的必要性。')

# 附录B 进度摘要表
t10=doc.tables[10]
prog=[
 ['第一阶段',
  '1.完成环境搭建：Vagrant+VirtualBox双虚拟机配置，网络参数(100Mbps/20ms)验证通过；2.完成基线运行：框架代码编译通过(make all)，client/server可通过UDP交换数据；3.完成框架代码架构分析：kernel仿真层、tju_packet报文层、tju_tcp协议层的模块职责、接口关系、数据流、并发机制(双线程+两把锁+条件变量)全部梳理清楚；4.完成环境搭建测试：4项测试(虚拟机连通性/编译/基线运行/UDP抓包)全部通过，记录了2个问题及修复方案；5.完成AI架构分析结论的人工核验：3项结论通过代码定位和运行验证，1项AI错误(tju_recv忙等误判)已修正；6.完成协议总体设计：报文格式(20字节头逐字段)、状态机FSM(11状态)、发送/接收缓冲区与滑动窗口协同、并发控制(三线程+两把锁+三个条件变量)、定时器(RTO/RTT估计/Karn/零窗口探测/TIME_WAIT)、7种异常处理全部设计完成；7.完成分阶段实现与测试策略规划(三阶段表格)；8.完成第一阶段实验报告撰写(基于V3模板)',
  '1.Git仓库地址：github.com/niuzeyu666/netproject(需执行git init和基线提交后推送)；2.数据结构扩展(sender_window_t/receiver_window_t/重传队列/失序队列)仅完成设计，代码中尚未实际启用；3.定时器线程仅完成设计，尚未实现；4.连接管理、可靠传输、流量控制、拥塞控制均未实现(属于第二、三阶段任务)；5.性能测试和trace分析尚未开展(第三阶段任务)；6.报告中截图位置为占位说明，需用户自行插入基线运行截图',
  '1.初始化Git仓库，提交基线代码和第一阶段设计文档，推送到github.com/niuzeyu666/netproject；2.第二阶段：实现连接管理(三次握手/四次挥手)、可靠传输(序号/ACK/超时重传/RTT估计)、流量控制(滑动窗口/零窗口)；3.第二阶段重点关注并发安全(条件变量替换忙等)和状态机正确性，充分测试丢包/乱序/重复场景；4.第二阶段实现时启用扩展的数据结构(sender_window_t/receiver_window_t/重传队列/失序队列)；5.第二阶段实现定时器线程，先实现SYN重传定时器，再扩展为数据重传和RTT估计'],
 ['第二阶段','【待第二阶段完成后填写】','【待第二阶段完成后填写】','【待第二阶段完成后填写】'],
 ['第三阶段','【待第三阶段完成后填写】','【待第三阶段完成后填写】','【第三阶段填写总结】'],
]
for i,r in enumerate(prog):
    for j,v in enumerate(r): sc(t10.rows[i+1].cells[j],v)

# 诚信声明日期
for p in doc.paragraphs:
    if '学生签名' in p.text:
        rp(p,'学生签名：____________________        日期：2026年09月06日')
        break

# 清理第一阶段范围内的提示段落
hints=['给出目录/模块表','至少包含一张主要函数调用图','分析tju_tcp_t、报文、缓冲区',
       '报告环境搭建测试结果','对AI架构分析结论的人工核验','给出总体架构图，说明连接管理',
       '说明课程自定义20字节报文头','说明连接状态、发送/接收缓冲区','说明线程/锁/条件变量、重传计时',
       '对AI工具与使用目的、AI建议摘要','至少选择一项对结果有实际影响的案例',
       '提供基线程序（初始框架）运行成功的截图']
for kw in hints:
    p=fp(kw)
    if p and '填写提示' in p.text:
        p._element.getparent().remove(p._element)

doc.save(DST)
print('附录和清理完成')
